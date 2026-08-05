# -*- coding: utf-8 -*-
import base64
import logging
import os
import subprocess
import tempfile

from odoo import _, api, fields, models
from odoo.exceptions import UserError

_logger = logging.getLogger(__name__)


class MembershipLetterWizard(models.TransientModel):
    _name = 'membership.letter.wizard'
    _description = 'Membership Letter Wizard'

    membership_id = fields.Many2one(
        'membership.membership',
        string='Membership',
        required=True,
    )
    template_id = fields.Many2one(
        'membership.letter.template',
        string='Letter Template',
        required=True,
    )
    output_format = fields.Selection(
        [('docx', 'Word Document (.docx)'), ('pdf', 'PDF')],
        string='Output Format',
        default='docx',
        required=True,
    )

    def _get_merge_values(self, membership):
        """Return a dict of merge field values for the given membership."""
        def _fmt_date(d):
            return d.strftime('%d/%m/%Y') if d else ''

        return {
            'member_number': membership.name or '',
            'member_name': membership.partner_name or '',
            'product_name': membership.product_id.name if membership.product_id else '',
            'status_name': membership.status_id.name if membership.status_id else '',
            'date_issued': _fmt_date(membership.date_issued),
            'date_expiry': _fmt_date(membership.date_expiry),
            'phone': membership.phone or '',
            'email': membership.email or '',
            'street': membership.street or '',
            'city': membership.city or '',
            'country': membership.country_id.name if membership.country_id else '',
        }

    def action_generate(self):
        self.ensure_one()

        if not self.template_id.document:
            raise UserError(_('The selected template has no document uploaded.'))

        # Try to import python-docx
        try:
            from docx import Document
        except ImportError:
            raise UserError(_(
                'The python-docx library is not installed. '
                'Please install it: pip install python-docx'
            ))

        membership = self.membership_id
        merge_values = self._get_merge_values(membership)

        # Decode the template document
        doc_bytes = base64.b64decode(self.template_id.document)
        doc_io = __import__('io').BytesIO(doc_bytes)

        try:
            doc = Document(doc_io)
        except Exception as e:
            raise UserError(_('Could not open the template document: %s') % str(e))

        # Replace {{field_name}} in all paragraphs and table cells
        def _replace_in_paragraph(para):
            for key, value in merge_values.items():
                placeholder = '{{' + key + '}}'
                if placeholder in para.text:
                    # Replace across runs (merge field may span multiple runs)
                    full_text = para.text
                    if placeholder in full_text:
                        new_text = full_text.replace(placeholder, value)
                        # Clear existing runs and set first run text
                        if para.runs:
                            para.runs[0].text = new_text
                            for run in para.runs[1:]:
                                run.text = ''
                        else:
                            para.add_run(new_text)

        # Process paragraphs in body
        for para in doc.paragraphs:
            _replace_in_paragraph(para)

        # Process tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        _replace_in_paragraph(para)

        # Save to BytesIO
        import io
        out_io = io.BytesIO()
        doc.save(out_io)
        docx_bytes = out_io.getvalue()

        filename_base = f'Member Letter - {membership.name}'

        if self.output_format == 'docx':
            # Return as downloadable attachment
            attachment = self.env['ir.attachment'].create({
                'name': f'{filename_base}.docx',
                'type': 'binary',
                'datas': base64.b64encode(docx_bytes),
                'mimetype': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                'res_model': 'membership.membership',
                'res_id': membership.id,
            })
            return {
                'type': 'ir.actions.act_url',
                'url': f'/web/content/{attachment.id}?download=true',
                'target': 'self',
            }

        elif self.output_format == 'pdf':
            # Write docx to tmp file and convert via LibreOffice
            try:
                with tempfile.TemporaryDirectory() as tmpdir:
                    docx_path = os.path.join(tmpdir, f'{filename_base}.docx')
                    with open(docx_path, 'wb') as f:
                        f.write(docx_bytes)

                    result = subprocess.run(
                        ['libreoffice', '--headless', '--convert-to', 'pdf',
                         '--outdir', tmpdir, docx_path],
                        capture_output=True,
                        timeout=60,
                    )

                    pdf_path = os.path.join(tmpdir, f'{filename_base}.pdf')
                    if result.returncode != 0 or not os.path.exists(pdf_path):
                        raise UserError(_(
                            'LibreOffice conversion failed. '
                            'Please install LibreOffice or use DOCX format instead.\n\n%s'
                        ) % (result.stderr.decode() or result.stdout.decode()))

                    with open(pdf_path, 'rb') as f:
                        pdf_bytes = f.read()

            except FileNotFoundError:
                raise UserError(_(
                    'LibreOffice is not installed or not in PATH. '
                    'Please install it or use DOCX output format.'
                ))
            except subprocess.TimeoutExpired:
                raise UserError(_('LibreOffice conversion timed out. Please try again.'))

            attachment = self.env['ir.attachment'].create({
                'name': f'{filename_base}.pdf',
                'type': 'binary',
                'datas': base64.b64encode(pdf_bytes),
                'mimetype': 'application/pdf',
                'res_model': 'membership.membership',
                'res_id': membership.id,
            })
            return {
                'type': 'ir.actions.act_url',
                'url': f'/web/content/{attachment.id}?download=true',
                'target': 'self',
            }
